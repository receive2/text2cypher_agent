#!/usr/bin/env python3
"""
auto_verify.py
==============
Automatically assess each needs-verification row by:
  1. Running a DuckDuckGo search for the from→to pair
  2. Asking GPT-4.1 to assign low/medium/high confidence + a reason

Checkpoints every row to benchmarks/verification_progress.jsonl so the run
is fully resume-safe.  After all rows are processed, writes the final
benchmarks/needs_verification_review.xlsx with new columns:
  auto_confidence  (low | medium | high)
  auto_reason      (evidence summary)
  evidence_urls    (search result URLs, pipe-separated)
  google_search    (clickable hyperlink — kept from original sheet)
  human_verification (dropdown, blank — for human override)
  notes            (blank — for human comments)

Usage
-----
  python scripts/auto_verify.py                  # run all rows
  python scripts/auto_verify.py --limit 20       # smoke-test first 20
  python scripts/auto_verify.py --reset          # wipe checkpoint and restart
"""

from __future__ import annotations

import argparse
import json
import re
import time
import urllib.parse
import warnings
from pathlib import Path

import requests

warnings.filterwarnings("ignore")  # suppress SSL verify=False warnings

# ── Paths ──────────────────────────────────────────────────────────────────────
REPO_ROOT    = Path(__file__).resolve().parent.parent
BENCHMARKS   = REPO_ROOT / "benchmarks"
CHECKPOINT   = BENCHMARKS / "verification_progress.jsonl"
OUTPUT_XLSX  = BENCHMARKS / "needs_verification_review.xlsx"

DATASETS = [
    "cypherbench_augmented_v2",
    "mindthequery_augmented_v2",
    "zograscope_augmented_v2",
]

# ── Import LLM helper ─────────────────────────────────────────────────────────
import sys
sys.path.insert(0, str(REPO_ROOT / "scripts"))
from walmart_llm import ask_llm  # noqa: E402


# ── DuckDuckGo HTML search ────────────────────────────────────────────────────
_DDG_URL     = "https://html.duckduckgo.com/html/"
_RESULT_RE   = re.compile(r'class="result__title".*?<a[^>]*>(.*?)</a>.*?class="result__snippet">(.*?)</a>', re.DOTALL)
_URL_RE      = re.compile(r'result__url[^>]*>(.*?)<', re.DOTALL)
_STRIP_TAGS  = re.compile(r'<[^>]+>')
_HEADERS     = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}


def _ddg_search(query: str, max_results: int = 4) -> list[dict]:
    """Return up to max_results hits [{title, snippet, url}].  Returns [] on error."""
    try:
        resp = requests.post(
            _DDG_URL,
            data={"q": query},
            headers=_HEADERS,
            timeout=15,
            verify=False,
        )
        if resp.status_code != 200:
            return []

        html   = resp.text
        titles = re.findall(r'class="result__title".*?<a[^>]*>(.*?)</a>', html, re.DOTALL)
        snips  = re.findall(r'class="result__snippet">(.*?)</a>',          html, re.DOTALL)
        urls   = re.findall(r'result__url[^>]*>\s*(.*?)\s*<',              html, re.DOTALL)

        results = []
        for t, s, u in zip(titles, snips, urls):
            results.append({
                "title":   _STRIP_TAGS.sub("", t).strip(),
                "snippet": _STRIP_TAGS.sub("", s).strip(),
                "url":     u.strip(),
            })
            if len(results) >= max_results:
                break
        return results
    except Exception:
        return []


# ── LLM assessment ────────────────────────────────────────────────────────────

_SYSTEM = """\
You are a linguistic fact-checker verifying whether a text rewrite is valid.
You will be given:
  - An original entity string ("from")
  - A rewritten entity string ("to")
  - The rewrite strategy (abbrev | alias | partial)
  - The graph domain (e.g. nba, flight_accident)
  - DuckDuckGo search results (may be empty)

Confidence rubric:
  high   = multiple independent sources confirm that "to" is a standard
           abbreviation/alias/partial-name for "from" in this domain, OR
           it is universally well-known (e.g. "CHI" = Chicago Bulls in NBA)
  medium = plausible and likely correct but ambiguous; only one weak source,
           or the rewrite is domain-plausible without strong confirmation
  low    = no evidence found, contradicted by search results, or the rewrite
           is implausible/wrong

Respond with ONLY valid JSON (no markdown, no extra text):
{
  "confidence": "low" | "medium" | "high",
  "reason": "<1-2 sentence evidence summary — apply to ALL confidence levels>"
}
"""


def _llm_assess(
    from_: str,
    to: str,
    strategy: str,
    graph: str,
    search_results: list[dict],
) -> dict:
    """Return {confidence, reason} from the LLM."""
    if search_results:
        snippets_txt = "\n".join(
            f"- [{r['title']}] {r['snippet']}  ({r['url']})"
            for r in search_results
        )
    else:
        snippets_txt = "(no search results found)"

    prompt = f"""\
from    : {from_}
to      : {to}
strategy: {strategy}
domain  : {graph}

Search results:
{snippets_txt}

Is "{to}" a valid {strategy} for "{from_}" in the {graph} domain?
"""

    for attempt in range(3):
        try:
            raw = ask_llm(prompt, system=_SYSTEM, max_tokens=256, temperature=0.0)
            # Strip potential markdown fences
            raw = raw.strip().lstrip("```json").lstrip("```").rstrip("```").strip()
            data = json.loads(raw)
            if data.get("confidence") in ("low", "medium", "high") and "reason" in data:
                return data
        except (json.JSONDecodeError, KeyError, RuntimeError):
            if attempt < 2:
                time.sleep(2)

    return {"confidence": "low", "reason": "LLM assessment failed after 3 attempts."}


# ── Data loading ──────────────────────────────────────────────────────────────

def _load_original_nl() -> dict[str, str]:
    mapping: dict[str, str] = {}
    for ds in DATASETS:
        rows = json.loads((BENCHMARKS / ds / "test.json").read_text())
        for r in rows:
            mapping[r["id"]] = r.get("_aug_meta", {}).get("original_nl") or r["nl"]
    return mapping


def _load_rows(orig_nl: dict[str, str]) -> list[dict]:
    rows = []
    for ds in DATASETS:
        with (BENCHMARKS / ds / "needs_verification.jsonl").open() as fh:
            for line in fh:
                r = json.loads(line)
                rows.append({
                    "dataset":     ds,
                    "graph":       r.get("graph", ""),
                    "qid":         r.get("qid", ""),
                    "strategy":    r.get("strategy", ""),
                    "original_nl": orig_nl.get(r.get("qid", ""), ""),
                    "augmented_nl": r.get("nl", ""),
                    "from":        r.get("from", ""),
                    "to":          r.get("to", ""),
                })
    return rows


def _load_checkpoint() -> dict[str, dict]:
    """Return {qid: result_dict} from the checkpoint file."""
    done: dict[str, dict] = {}
    if CHECKPOINT.exists():
        with CHECKPOINT.open() as fh:
            for line in fh:
                obj = json.loads(line)
                done[obj["qid"]] = obj
    return done


# ── Excel generation ──────────────────────────────────────────────────────────

def _build_excel(rows: list[dict], results: dict[str, dict]) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
    EVEN_FILL    = PatternFill("solid", fgColor="EBF3FB")
    ODD_FILL     = PatternFill("solid", fgColor="FFFFFF")
    CONF_COLORS  = {
        "high":   PatternFill("solid", fgColor="C6EFCE"),  # green
        "medium": PatternFill("solid", fgColor="FFEB9C"),  # yellow
        "low":    PatternFill("solid", fgColor="FFC7CE"),  # red
    }
    WRAP   = Alignment(wrap_text=True, vertical="top")
    CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
    BORDER = Border(
        bottom=Side(style="thin", color="BDD7EE"),
        right=Side(style="thin", color="BDD7EE"),
    )

    COLUMNS = [
        ("dataset",             22),
        ("graph",               16),
        ("qid",                 36),
        ("strategy",            12),
        ("original_nl",         52),
        ("augmented_nl",        52),
        ("from",                28),
        ("to",                  24),
        ("auto_confidence",     18),
        ("auto_reason",         60),
        ("evidence_urls",       40),
        ("google_search",       18),
        ("human_verification",  22),
        ("notes",               40),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verification"

    headers = [c for c, _ in COLUMNS]
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 24

    CONF_COL   = headers.index("auto_confidence")    + 1
    GOOGLE_COL = headers.index("google_search")      + 1
    VERIFY_COL = headers.index("human_verification") + 1

    for row_idx, r in enumerate(rows, start=2):
        base_fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL
        res  = results.get(r["qid"], {})
        conf = res.get("confidence", "")
        urls = res.get("evidence_urls", [])

        values = [
            r["dataset"],
            r["graph"],
            r["qid"],
            r["strategy"],
            r["original_nl"],
            r["augmented_nl"],
            r["from"],
            r["to"],
            conf,
            res.get("reason", ""),
            " | ".join(urls),
            "🔍 Search",
            "",
            "",
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill      = base_fill
            cell.border    = BORDER
            cell.alignment = WRAP

        # Confidence cell — colour-coded
        conf_cell = ws.cell(row=row_idx, column=CONF_COL)
        if conf in CONF_COLORS:
            conf_cell.fill = CONF_COLORS[conf]
        conf_cell.alignment = CENTER

        # Google search hyperlink
        q    = urllib.parse.quote_plus(f'"{r["from"]}" "{r["to"]}" {r["graph"]}')
        url  = f"https://www.google.com/search?q={q}"
        gcell = ws.cell(row=row_idx, column=GOOGLE_COL)
        gcell.hyperlink = url
        gcell.value     = "🔍 Search"
        gcell.font      = Font(color="0563C1", underline="single")
        gcell.alignment = CENTER

        ws.row_dimensions[row_idx].height = 52

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    dv = DataValidation(
        type="list",
        formula1='"CORRECT,INCORRECT,UNSURE"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Invalid value",
        error="Choose CORRECT, INCORRECT, or UNSURE.",
    )
    ws.add_data_validation(dv)
    last = len(rows) + 1
    dv.add(f"{get_column_letter(VERIFY_COL)}2:{get_column_letter(VERIFY_COL)}{last}")

    wb.save(OUTPUT_XLSX)
    print(f"Excel written → {OUTPUT_XLSX}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(description="Auto-verify needs_verification rows.")
    parser.add_argument("--limit", type=int, default=0, help="Process only the first N rows (0 = all)")
    parser.add_argument("--reset", action="store_true",  help="Delete checkpoint and restart from scratch")
    parser.add_argument("--excel-only", action="store_true", help="Skip LLM/search; rebuild Excel from checkpoint")
    parser.add_argument("--delay", type=float, default=2.5, help="Seconds between DDG searches (default 2.5)")
    args = parser.parse_args()

    if args.reset and CHECKPOINT.exists():
        CHECKPOINT.unlink()
        print("Checkpoint cleared.")

    print("Loading data …")
    orig_nl = _load_original_nl()
    rows    = _load_rows(orig_nl)
    if args.limit:
        rows = rows[: args.limit]

    done = _load_checkpoint()
    print(f"  {len(rows)} rows total — {len(done)} already done, {len(rows) - len(done)} remaining")

    if not args.excel_only:
        todo = [r for r in rows if r["qid"] not in done]

        with CHECKPOINT.open("a") as cp:
            for i, r in enumerate(todo, start=1):
                qid = r["qid"]
                print(f"[{i}/{len(todo)}] {r['from']!r} → {r['to']!r}  ({r['strategy']}, {r['graph']})", end=" … ", flush=True)

                # 1. Search
                query   = f'"{r["from"]}" "{r["to"]}" {r["graph"]}'
                hits    = _ddg_search(query)
                urls    = [h["url"] for h in hits]

                # 2. LLM assess
                assessment = _llm_assess(
                    from_=r["from"],
                    to=r["to"],
                    strategy=r["strategy"],
                    graph=r["graph"],
                    search_results=hits,
                )

                result = {
                    "qid":           qid,
                    "confidence":    assessment["confidence"],
                    "reason":        assessment["reason"],
                    "evidence_urls": urls,
                }
                done[qid] = result
                cp.write(json.dumps(result) + "\n")
                cp.flush()

                print(f"{assessment['confidence'].upper()}")

                # Rate-limit DDG
                time.sleep(args.delay)

    print("\nBuilding Excel …")
    _build_excel(rows, done)
    print("Done.")


if __name__ == "__main__":
    main()
