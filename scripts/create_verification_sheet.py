#!/usr/bin/env python3
"""
create_verification_sheet.py
=============================
Generates benchmarks/needs_verification_review.xlsx for human review of
augmented NL ↔ gold-Cypher pairs across all three benchmark datasets.

Columns
-------
dataset          : source benchmark
graph            : Neo4j graph name
qid              : unique question id
strategy         : abbrev | alias | partial
original_nl      : natural-language before augmentation
augmented_nl     : natural-language after augmentation (what the agent sees)
from             : original entity string
to               : rewritten entity string
google_search    : clickable hyperlink → Google search for from/to validity
human_verification: CORRECT | INCORRECT | UNSURE  (dropdown, blank to start)
notes            : freeform reviewer notes
"""

from __future__ import annotations

import json
import urllib.parse
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ── Paths ──────────────────────────────────────────────────────────────────────
REPO_ROOT   = Path(__file__).resolve().parent.parent
BENCHMARKS  = REPO_ROOT / "benchmarks"
OUTPUT_PATH = BENCHMARKS / "needs_verification_review.xlsx"

DATASETS = [
    "cypherbench_augmented_v2",
    "mindthequery_augmented_v2",
    "zograscope_augmented_v2",
]

# ── Helpers ────────────────────────────────────────────────────────────────────

def google_search_url(from_: str, to: str, graph: str) -> str:
    """Build a Google search URL to help verify whether `from` → `to` is valid."""
    query = f'"{from_}" "{to}" {graph}'
    return "https://www.google.com/search?q=" + urllib.parse.quote_plus(query)


def load_original_nl() -> dict[str, str]:
    """Return qid → original_nl from every test.json (before augmentation)."""
    mapping: dict[str, str] = {}
    for ds in DATASETS:
        path = BENCHMARKS / ds / "test.json"
        rows = json.loads(path.read_text())
        for r in rows:
            aug_meta = r.get("_aug_meta", {})
            mapping[r["id"]] = aug_meta.get("original_nl") or r["nl"]
    return mapping


def load_rows(orig_nl: dict[str, str]) -> list[dict]:
    """Merge needs_verification.jsonl from all datasets."""
    rows = []
    for ds in DATASETS:
        path = BENCHMARKS / ds / "needs_verification.jsonl"
        with path.open() as fh:
            for line in fh:
                r = json.loads(line)
                rows.append(
                    {
                        "dataset":    ds,
                        "graph":      r.get("graph", ""),
                        "qid":        r.get("qid", ""),
                        "strategy":   r.get("strategy", ""),
                        "original_nl":  orig_nl.get(r.get("qid", ""), ""),
                        "augmented_nl": r.get("nl", ""),
                        "from":       r.get("from", ""),
                        "to":         r.get("to", ""),
                    }
                )
    return rows


# ── Styles ─────────────────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")   # dark blue
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
EVEN_FILL    = PatternFill("solid", fgColor="EBF3FB")   # light blue tint
ODD_FILL     = PatternFill("solid", fgColor="FFFFFF")
WRAP_ALIGN   = Alignment(wrap_text=True, vertical="top")
CENTER_ALIGN = Alignment(horizontal="center", vertical="top", wrap_text=True)
THIN_BORDER  = Border(
    bottom=Side(style="thin", color="BDD7EE"),
    right=Side(style="thin", color="BDD7EE"),
)

COLUMNS = [
    ("dataset",             22),
    ("graph",               18),
    ("qid",                 36),
    ("strategy",            12),
    ("original_nl",         52),
    ("augmented_nl",        52),
    ("from",                28),
    ("to",                  28),
    ("google_search",       20),
    ("human_verification",  22),
    ("notes",               40),
]


# ── Main ───────────────────────────────────────────────────────────────────────

def build_workbook(rows: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verification"

    # ── Header row ────────────────────────────────────────────────────────────
    headers = [col for col, _ in COLUMNS]
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 24

    # ── Data rows ─────────────────────────────────────────────────────────────
    GOOGLE_COL  = headers.index("google_search") + 1
    VERIFY_COL  = headers.index("human_verification") + 1

    for row_idx, r in enumerate(rows, start=2):
        fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL

        values = [
            r["dataset"],
            r["graph"],
            r["qid"],
            r["strategy"],
            r["original_nl"],
            r["augmented_nl"],
            r["from"],
            r["to"],
            "🔍 Search",        # placeholder; hyperlink set below
            "",                 # human_verification — blank
            "",                 # notes
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill      = fill
            cell.border    = THIN_BORDER
            cell.alignment = WRAP_ALIGN

        # ── Google search hyperlink ───────────────────────────────────────────
        url  = google_search_url(r["from"], r["to"], r["graph"])
        gcell = ws.cell(row=row_idx, column=GOOGLE_COL)
        gcell.hyperlink = url
        gcell.value     = "🔍 Search"
        gcell.font      = Font(color="0563C1", underline="single")
        gcell.alignment = CENTER_ALIGN

        # ── Row height ────────────────────────────────────────────────────────
        ws.row_dimensions[row_idx].height = 52

    # ── Freeze header + auto-filter ───────────────────────────────────────────
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ── Data validation dropdown on human_verification ────────────────────────
    dv = DataValidation(
        type="list",
        formula1='"CORRECT,INCORRECT,UNSURE"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Invalid value",
        error="Choose CORRECT, INCORRECT, or UNSURE from the dropdown.",
    )
    ws.add_data_validation(dv)
    last_row = len(rows) + 1
    dv.add(f"{get_column_letter(VERIFY_COL)}2:{get_column_letter(VERIFY_COL)}{last_row}")

    return wb


def main() -> None:
    print("Loading original NL mappings …")
    orig_nl = load_original_nl()

    print("Loading needs_verification rows …")
    rows = load_rows(orig_nl)
    print(f"  → {len(rows)} rows across {len(DATASETS)} datasets")

    print("Building workbook …")
    wb = build_workbook(rows)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Saved → {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
